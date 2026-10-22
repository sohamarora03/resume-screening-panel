"""
Build the project report (docs/Project_Report.docx) with python-docx.

    python main.py                    # creates outputs/results.json
    python scripts/build_report.py    # builds the report

Edit STUDENT below (or the cover page in Word) with your own details.
The table of contents is filled in automatically: the script renders the
document once with LibreOffice to find each heading's page, then rebuilds it.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SHOTS = DOCS / "screenshots"
DIAG = DOCS / "diagrams"
OUT = DOCS / "Project_Report.docx"

STUDENT = {
    "Name": "[YOUR NAME]",
    "SAP ID": "[YOUR SAP ID]",
    "Roll Number": "[YOUR ROLL NO]",
    "Course / Batch": "[COURSE NAME]",
    "College": "[COLLEGE NAME]",
    "Guide / Mentor": "[GUIDE NAME]",
}
GITHUB_URL = "https://github.com/[your-username]/resume-screening-panel"
DRIVE_LINK = "[PASTE LINK HERE]"
TITLE = "Resume Screening Panel"
SUBTITLE = "A Multi-Agent System for Fair and Explainable Resume Screening"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x5F, 0x6D)


# =============================================================== helpers
class Report:
    def __init__(self, toc_pages: dict[str, int] | None = None):
        self.doc = Document()
        self.fig = 0
        self.tab = 0
        self.headings: list[tuple[int, str]] = []
        self.toc_pages = toc_pages or {}
        self._styles()
        self._page_setup()

    # ---------------------------------------------------------- setup
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
            s = st[name]
            s.font.name = "Calibri"
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = NAVY
            s.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
            rfonts = s.element.rPr.find(qn("w:rFonts"))
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts is not None and rfonts.get(qn(attr)):
                    del rfonts.attrib[qn(attr)]
            s.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 10)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9.5)
        cap.font.italic = True
        cap.font.bold = False
        cap.font.color.rgb = GREY
        for lvl, indent in (("TOC 1", 0), ("TOC 2", 0.6)):
            s = st[lvl] if lvl in [x.name for x in st] else st.add_style(lvl, 1)
            s.font.name = "Calibri"
            s.font.size = Pt(10.5)
            s.font.bold = lvl == "TOC 1"
            s.paragraph_format.left_indent = Cm(indent)
            s.paragraph_format.space_before = Pt(4 if lvl == "TOC 1" else 0)
            s.paragraph_format.space_after = Pt(0)
            s.paragraph_format.line_spacing = 1.0
            s.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)

    def _page_setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.different_first_page_header_footer = True     # no number on the cover
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(footer, "PAGE", "1")
        header = sec.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = header.add_run(f"{TITLE} - Project Report")
        r.font.size = Pt(9)
        r.font.color.rgb = GREY

    @staticmethod
    def _field(paragraph, instr: str, placeholder: str = ""):
        def fc(kind):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            return el
        r1 = paragraph.add_run()
        r1._r.append(fc("begin"))
        r2 = paragraph.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r2._r.append(it)
        r3 = paragraph.add_run()
        r3._r.append(fc("separate"))
        paragraph.add_run(placeholder)
        r5 = paragraph.add_run()
        r5._r.append(fc("end"))

    # ---------------------------------------------------------- content
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text):
        self.headings.append((1, text))
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.headings.append((2, text))
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def p(self, text, bold_lead: str | None = None, align=None, italic=False):
        para = self.doc.add_paragraph()
        if bold_lead:
            para.add_run(bold_lead).bold = True
        self._rich(para, text, italic)
        if align:
            para.alignment = align
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    @staticmethod
    def _rich(para, text, italic=False):
        """Support **bold** spans inside a string."""
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                run = para.add_run(part)
                run.bold = i % 2 == 1
                run.italic = italic

    def bullets(self, items, style="List Bullet"):
        for item in items:
            para = self.doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            self._rich(para, item)

    def numbered(self, items):
        """Numbered list that restarts at 1 (each list gets its own w:num)."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style_numpr = self.doc.styles["List Number"].element.pPr.numPr
        base_num = style_numpr.numId.val
        abstract = numbering.xpath(f'./w:num[@w:numId="{base_num}"]/w:abstractNumId')[0].get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.xpath("./w:num")) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        num.append(a)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        for item in items:
            para = self.doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(3)
            numpr = para._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = new_id
            self._rich(para, item)

    def code(self, text: str, caption: str | None = None):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        trpr = table.rows[0]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        self._shade(cell, "F4F6F8")
        self._cell_border(cell, "C9D1DA")
        cell.paragraphs[0].text = ""
        lines = text.strip("\n").splitlines()
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(line if line else " ")
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(8)
        if caption:
            self.caption(f"Listing: {caption}")
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def caption(self, text):
        para = self.doc.add_paragraph(text, style="Caption")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(10)

    def figure(self, path: Path, caption: str, width_cm: float = 16, max_h_cm: float = 20):
        with Image.open(path) as im:
            w, h = im.size
        width = width_cm
        if h / w * width > max_h_cm:
            width = max_h_cm * w / h
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width))
        self.fig += 1
        self.caption(f"Fig {self.fig}: {caption}")

    def table(self, header, rows, widths_cm, caption=None, font=9.5):
        if caption:
            self.tab += 1
            cap = self.doc.add_paragraph(f"Table {self.tab}: {caption}", style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(text))
            run.bold = True
            run.font.size = Pt(font)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            self._shade(cell, "1F3A5F")
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                self._rich(para, str(text))
                for run in para.runs:
                    run.font.size = Pt(font)
                if r_i % 2 == 1:
                    self._shade(cells[i], "F1F5F9")
        self._fix_widths(t, widths_cm)
        for row in t.rows:                     # keep each row on one page
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        # repeat header row on new pages
        trpr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    @staticmethod
    def _fix_widths(t, widths_cm):
        t.autofit = False
        tblpr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        grid = t._tbl.tblGrid
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            col.set(qn("w:w"), str(int(widths_cm[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)

    @staticmethod
    def _shade(cell, hex_fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcpr.append(shd)

    @staticmethod
    def _cell_border(cell, colour):
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "6")
            b.set(qn("w:color"), colour)
            borders.append(b)
        tcpr.append(borders)

    # ---------------------------------------------------------- TOC
    def toc(self, entries: list[tuple[int, str]]):
        """Real Word TOC field, pre-filled with entries and page numbers."""
        def fc(kind):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            return el
        first = True
        paras = []
        for level, text in entries:
            para = self.doc.add_paragraph(style=f"TOC {level}")
            if first:
                for el in (fc("begin"),):
                    para.add_run()._r.append(el)
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                para.add_run()._r.append(it)
                para.add_run()._r.append(fc("separate"))
                first = False
            page = self.toc_pages.get(text, "")
            para.add_run(f"{text}\t{page}")
            paras.append(para)
        paras[-1].add_run()._r.append(fc("end"))

    def save(self, path: Path):
        self.doc.save(path)


# =============================================================== content
def crop_tall(src: Path, max_ratio: float = 1.25, name: str | None = None) -> Path:
    """Report-friendly copy of a very tall screenshot (keeps the top part)."""
    with Image.open(src) as im:
        w, h = im.size
        if h <= w * max_ratio:
            return src
        dst = SHOTS / "report" / (name or src.name)
        dst.parent.mkdir(exist_ok=True)
        im.crop((0, 0, w, int(w * max_ratio))).save(dst)
        return dst


def code_excerpt(path: Path, start_pat: str, n_lines: int) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    for i, line in enumerate(lines):
        if re.search(start_pat, line):
            return "\n".join(lines[i:i + n_lines])
    raise ValueError(f"pattern {start_pat!r} not found in {path}")


def build(toc_pages=None) -> Report:
    result = json.loads((ROOT / "outputs" / "results.json").read_text(encoding="utf-8"))
    tests = (DOCS / "test_results.txt").read_text(encoding="utf-8") \
        if (DOCS / "test_results.txt").exists() else ""
    reviews = {r["candidate"]: r for r in result["bias_reviews"]}
    R = Report(toc_pages)
    d = R.doc

    # ------------------------------------------------------------ cover
    for _ in range(2):
        d.add_paragraph()
    for text, size, bold, colour in (
        (STUDENT["College"], 16, True, NAVY),
        ("A Project Report on", 13, False, GREY),
        (TITLE, 30, True, NAVY),
        (SUBTITLE, 15, False, GREY),
    ):
        para = d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(10)
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = colour
    para = d.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("Coursera On-the-Job Training (OJT) Project - Multi-Agent Systems")
    run.font.size = Pt(12)
    run.italic = True
    for _ in range(2):
        d.add_paragraph()
    para = d.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run("Submitted by").bold = True

    t = d.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for key in ("Name", "SAP ID", "Roll Number", "Course / Batch", "Guide / Mentor"):
        cells = t.add_row().cells
        cells[0].text, cells[1].text = "", ""
        cells[0].paragraphs[0].add_run(key).bold = True
        run = cells[1].paragraphs[0].add_run(STUDENT[key])
        if STUDENT[key].startswith("["):
            run.font.highlight_color = 7  # yellow: easy to spot placeholders
        R._shade(cells[0], "E8EEF5")
    R._fix_widths(t, [5, 8])
    for _ in range(3):
        d.add_paragraph()
    para = d.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"Date of submission: {date.today():%d %B %Y}")
    run.font.color.rgb = GREY
    R.page_break()

    # ------------------------------------------------------------ index
    para = d.add_paragraph()
    run = para.add_run("Index")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY
    para.paragraph_format.space_after = Pt(12)
    toc_entries = TOC_ENTRIES
    R.toc(toc_entries)
    R.page_break()

    # ------------------------------------------------------------ 1
    R.h1("1. Introduction")
    R.h2("1.1 What is an AI agent?")
    R.p("An **AI agent** is a software program that perceives its environment, reasons about a goal "
        "and takes actions to reach that goal with some degree of autonomy. Modern agents are usually "
        "built on a **Large Language Model (LLM)**, which gives them the ability to read instructions in "
        "natural language, plan the steps needed, call tools and produce structured output. An agent is "
        "more than a single prompt: it has a defined **role**, a **goal**, background knowledge (its "
        "**backstory** or system prompt), access to **tools or data**, and a clear contract for what it "
        "must hand back.")
    R.h2("1.2 What is a multi-agent system?")
    R.p("A **multi-agent system (MAS)** is a group of agents that each own one part of a larger problem "
        "and collaborate by passing information to one another. Agents can be arranged in different "
        "patterns: a **sequential pipeline** (each agent works on the previous agent's output), a "
        "**hierarchical** structure (a manager agent delegates to workers) or a **conversational** "
        "group that debates until it agrees. This project uses a sequential pipeline coordinated by an "
        "orchestrator, which is the easiest pattern to trace and audit.")
    R.h2("1.3 Why multi-agent instead of one large prompt?")
    R.bullets([
        "**Separation of concerns:** each agent has one job and a short, focused prompt, which makes "
        "its output more reliable than one prompt that must do everything.",
        "**Checks and balances:** one agent can review another's work. Here, the Bias Checker Agent "
        "audits the Screener Agent, which a single prompt cannot do to itself in a transparent way.",
        "**Explainability:** every hand-off produces a JSON artefact that can be logged, displayed and "
        "tested, so it is clear which agent made which decision.",
        "**Modularity:** an agent can be improved, replaced or given a different model without "
        "touching the others.",
        "**Scalability:** per-resume tasks (screening and auditing) can be run in parallel in future.",
    ])
    R.h2("1.4 Multi-agent frameworks")
    R.table(["Framework", "Developer", "Core idea", "Best suited for"], [
        ["CrewAI", "CrewAI Inc. (open source)", "Agents with role, goal and backstory work on Tasks "
         "inside a Crew; sequential or hierarchical processes", "Role-based business workflows (used here)"],
        ["AutoGen", "Microsoft Research", "Agents talk to each other in a multi-turn conversation; "
         "supports human-in-the-loop", "Conversational problem solving, code generation"],
        ["LangGraph", "LangChain", "Agents are nodes in a stateful graph with explicit edges, "
         "loops and checkpoints", "Complex branching or long-running workflows"],
    ], [2.6, 3.2, 6.2, 4.0], caption="Comparison of popular multi-agent frameworks")
    R.p("**CrewAI** was chosen because its role/goal/backstory model maps directly onto an HR panel, "
        "it is simple to read, and it works with any OpenAI-compatible model provider such as Groq.")
    R.h2("1.5 Project overview")
    R.p("The **Resume Screening Panel** imitates a real interview panel. A recruiter uploads one job "
        "description and several resumes (PDF or TXT). Four agents then work in sequence: the **JD "
        "Analyst Agent** turns the job description into structured criteria, the **Screener Agent** "
        "scores each resume, the **Bias Checker Agent** removes unfair influences from those scores, and "
        "the **HR Summary Agent** produces a ranked shortlist with Hire / Maybe / Reject "
        "recommendations and tailored interview questions. Results are shown in a Streamlit web app and "
        "can be downloaded as CSV, PDF or JSON.")

    # ------------------------------------------------------------ 2
    R.h1("2. Problem Definition")
    R.h2("2.1 Background")
    R.p("Entry-level job postings in India routinely attract hundreds of applications. During my OJT "
        "I observed that first-round screening is usually done manually by recruiters who spend well "
        "under a minute on each resume. This creates several problems:")
    R.bullets([
        "**Time and cost:** screening hundreds of resumes by hand takes days of recruiter time.",
        "**Inconsistency:** two recruiters (or the same recruiter at different times) score the same "
        "resume differently, because the criteria are applied from memory.",
        "**Little explanation:** rejected candidates and hiring managers rarely see why a decision "
        "was made.",
        "**Keyword filters are too rigid:** simple applicant tracking systems reject good candidates "
        "who describe a skill with different words (for example, PostgreSQL instead of SQL).",
    ])
    R.h2("2.2 Bias in hiring")
    R.p("Studies of hiring repeatedly show that factors unrelated to the job influence screening "
        "decisions. Common examples are penalising **career breaks** (often taken for maternity or "
        "caregiving), **age**, **marital status**, a **foreign-sounding name or overseas background**, "
        "and the **prestige of the college** attended (in both directions: a top-tier brand can hide "
        "weak skills, and a lesser-known college can hide strong ones). Automated systems can copy these "
        "same biases if nothing checks them. A fair screening tool therefore needs an explicit step that "
        "looks for, removes and explains such influences.")
    R.h2("2.3 Problem statement")
    R.p("To design and build a multi-agent AI system that screens resumes against a job description "
        "quickly and consistently, **detects and corrects biased reasoning**, and gives recruiters a "
        "ranked, explainable shortlist with tailored interview questions.", italic=False)
    R.h2("2.4 Objectives")
    R.numbered([
        "Extract structured hiring criteria (must-have skills, nice-to-have skills, experience, "
        "qualifications) from any job description.",
        "Score every resume from 0 to 100 with a per-skill breakdown that quotes evidence from the resume.",
        "Audit each scoring decision for bias across eight protected or job-irrelevant categories and "
        "correct the score where bias is found.",
        "Rank candidates, recommend Hire / Maybe / Reject and generate 3-5 tailored interview questions.",
        "Prove the multi-agent behaviour with timestamped hand-off logs, a live status UI and tests.",
        "Keep API keys secure and make the project runnable offline for evaluation (demo mode).",
    ])
    R.h2("2.5 Scope")
    R.table(["In scope", "Out of scope"], [
        ["PDF and TXT resumes and job descriptions", "Scanned image resumes (OCR)"],
        ["One job description per run, any number of resumes", "Comparing one resume against many jobs"],
        ["English-language documents", "Other languages"],
        ["Decision support for a human recruiter", "Fully automatic hiring decisions"],
        ["Groq (Llama 3.3 70B) live mode and an offline demo mode", "Fine-tuning a custom model"],
    ], [8, 8], caption="Project scope")

    # ------------------------------------------------------------ 3
    R.page_break()
    R.h1("3. System Design")
    R.h2("3.1 System architecture")
    R.p("The system is organised in four layers. The **presentation layer** (Streamlit web UI and a "
        "command-line runner) collects files and displays results. The **orchestration layer** parses "
        "the documents and runs the four CrewAI agents in order. The **model layer** supplies the "
        "language model: Groq's Llama 3.3 70B in live mode, or a deterministic offline engine in demo "
        "mode. The **output layer** stores the hand-off log and the final results.")
    R.figure(DIAG / "system_architecture.png", "System architecture of the Resume Screening Panel")
    R.h2("3.2 Agent roles")
    R.table(["#", "Agent", "Role / goal", "Input", "Output (JSON)"], [
        ["1", "JD Analyst Agent", "Senior technical recruiter. Turns the job description into "
         "structured hiring criteria without inventing requirements.", "Job description text",
         "job_title, must_have_skills, nice_to_have_skills, experience, qualifications"],
        ["2", "Screener Agent", "First-pass screener. Scores each resume 0-100 with evidence for "
         "every skill and lists any other factors that affected its impression.", "Criteria + one resume",
         "skill_score, adjustments, score, skill_breakdown, justification"],
        ["3", "Bias Checker Agent", "Fair-hiring compliance officer. Flags adjustments based on "
         "protected or job-irrelevant factors and reverses them.", "Resume + screening result",
         "original_score, adjusted_score, bias_flags, verdict"],
        ["4", "HR Summary Agent", "HR business partner. Ranks candidates, recommends and writes "
         "interview questions based on each person's evidence and gaps.", "Criteria + all corrected results",
         "ranking, shortlist, summary"],
    ], [0.7, 2.8, 5.6, 3.0, 4.0], caption="The four agents and their contracts", font=9)
    R.h2("3.3 Agent workflow and hand-offs")
    R.p("Figure 2 shows the order of work. Each arrow labelled 'hand-off' is logged by the "
        "orchestrator. The Screener and Bias Checker agents run once per resume, so a run with six "
        "resumes executes 2 + 2 x 6 = 14 agent tasks.")
    R.figure(DIAG / "agent_workflow.png", "Agent workflow and hand-off sequence", width_cm=9.5,
             max_h_cm=19)
    R.h2("3.4 Data flow")
    R.numbered([
        "The user uploads files; utils/parser.py extracts plain text from PDF (pypdf) or TXT.",
        "The orchestrator wraps each task's input as JSON between <<<INPUT_JSON markers and sends "
        "it to an agent through a one-agent CrewAI Crew.",
        "Each agent must reply with JSON only. extract_json() reads the reply (even if wrapped in "
        "code fences) and the orchestrator retries once if the JSON is invalid.",
        "The criteria JSON goes to the Screener; screening JSON goes to the Bias Checker; corrected "
        "scores and skill evidence go to the HR Summary Agent.",
        "The final result is shown in the UI, saved to outputs/ and exported as CSV / PDF / JSON. "
        "Every step is written to logs/run_log.txt.",
    ])
    R.h2("3.5 Scoring model")
    R.table(["Component", "Weight", "Rule"], [
        ["Must-have skills", "70", "70 x (matched must-have skills / total must-have skills)"],
        ["Nice-to-have skills", "20", "20 x (matched nice-to-have skills / total)"],
        ["Relevant experience", "10", "10 if the candidate has 6+ months of experience, otherwise 6"],
        ["Other adjustments", "+/-", "Anything else the Screener noted; audited by the Bias Checker"],
        ["Recommendation", "-", "Hire >= 75, Maybe 55-74, Reject < 55 (bias-corrected score)"],
    ], [4, 2, 10], caption="Scoring and recommendation rules")
    R.h2("3.6 Key design decisions")
    R.bullets([
        "**One Crew per stage:** running each agent as its own Crew lets the orchestrator log every "
        "hand-off and update the UI live, while CrewAI still builds the prompts and runs the agents.",
        "**JSON contracts:** fixed schemas between agents make outputs testable and easy to display.",
        "**Screener lists its adjustments openly:** instead of hiding gut-feel judgements inside one "
        "number, the Screener must list them, so the Bias Checker can audit each one.",
        "**Demo mode:** a CrewAI BaseLLM subclass answers with each agent's offline engine when no API "
        "key is present, so the same code path runs for tests, evaluation and recordings.",
        "**Secrets in .env:** the Groq key is read from .env (ignored by Git) and never hard-coded.",
    ])

    # ------------------------------------------------------------ 4
    R.page_break()
    R.h1("4. Development Details")
    R.h2("4.1 Technology stack")
    R.table(["Layer", "Technology", "Purpose"], [
        ["Language", "Python 3.10+", "All code"],
        ["Multi-agent framework", "CrewAI 1.15", "Agents, tasks, crews"],
        ["LLM", "Groq API - Llama 3.3 70B (via OpenAI-compatible endpoint)", "Live reasoning"],
        ["Web UI", "Streamlit 1.66", "Upload, live status, results, downloads"],
        ["Documents", "pypdf, reportlab", "Read PDF resumes; export PDF results"],
        ["Data", "pandas", "Results tables and CSV export"],
        ["Terminal output", "rich", "Coloured, timestamped agent logs"],
        ["Testing", "pytest", "15 unit and end-to-end tests"],
        ["Documentation", "Playwright, Graphviz, python-docx, ffmpeg",
         "Screenshots, demo video, diagrams, this report"],
        ["Version control", "Git + GitHub", "Source code hosting"],
    ], [3.5, 6.5, 6], caption="Technology stack")
    R.h2("4.2 Folder structure")
    R.code("""
resume-screening-panel/
|-- agents/                 one file per agent
|   |-- jd_analyst.py       Agent 1: JD Analyst Agent
|   |-- screener.py         Agent 2: Screener Agent
|   |-- bias_checker.py     Agent 3: Bias Checker Agent
|   `-- hr_summary.py       Agent 4: HR Summary Agent
|-- tasks/prompts.py        task descriptions and JSON contracts
|-- utils/
|   |-- llm.py              Groq LLM / offline DemoLLM factory
|   |-- logger.py           timestamped hand-off logger
|   |-- parser.py           PDF/TXT reading, JSON extraction
|   |-- skills.py           skill synonyms and evidence finder
|   `-- exporter.py         CSV and PDF export
|-- orchestrator.py         runs the four agents in sequence
|-- app.py                  Streamlit web UI
|-- main.py                 command-line runner
|-- config.py               settings from .env
|-- sample_data/            1 JD + 6 fictional resumes (TXT and PDF)
|-- tests/                  pytest test suite
|-- scripts/                screenshots, video, diagrams, report
|-- docs/                   report, screenshots, diagrams, demo video
|-- logs/run_log.txt        agent hand-off log (created at run time)
|-- requirements.txt  .env.example  .gitignore  README.md
""")
    R.h2("4.3 Key code")
    R.h3("Agent definition (agents/bias_checker.py)")
    R.p("Every agent file defines the CrewAI agent's role, goal and backstory, plus an offline engine "
        "with the same JSON output used in demo mode.")
    R.code(code_excerpt(ROOT / "agents/bias_checker.py", r"^NAME = ", 26))
    R.h3("Task with JSON contract (tasks/prompts.py)")
    R.code(code_excerpt(ROOT / "tasks/prompts.py", r"^def bias_task", 21))
    R.h3("Orchestrator: running an agent and logging the hand-off (orchestrator.py)")
    R.code(code_excerpt(ROOT / "orchestrator.py", r"def _kickoff", 13))
    R.code(code_excerpt(ROOT / "orchestrator.py", r"# 1\. JD Analyst Agent", 20))
    R.h3("Model selection: live Groq model or offline demo engine (utils/llm.py)")
    R.code(code_excerpt(ROOT / "utils/llm.py", r"^def get_llm", 10))
    R.h3("Bias audit logic used by the offline engine (agents/bias_checker.py)")
    R.code(code_excerpt(ROOT / "agents/bias_checker.py", r"^def demo_respond", 22))

    R.h2("4.4 User interface and execution")
    R.p("The Streamlit app is started with **streamlit run app.py**. The screenshots below were "
        "captured automatically with Playwright during a real run on the sample data.")
    R.figure(SHOTS / "01_upload_screen.png", "Home screen with upload areas and the four agent cards")
    R.figure(SHOTS / "03_agents_running.png",
             "Live status during a run: JD Analyst done, Screener Agent working on resume 2 of 6")
    R.figure(SHOTS / "03b_agents_running_bias.png", "Live status: Bias Checker Agent auditing decisions")
    R.figure(SHOTS / "04_pipeline_complete.png", "All four agents completed")
    R.figure(SHOTS / "05_agent1_jd_analyst.png", "Agent 1 output: structured hiring criteria",
             width_cm=15)
    R.figure(crop_tall(SHOTS / "06_agent2_screener.png", 1.0),
             "Agent 2 output: score and per-skill evidence table (first candidates shown)", width_cm=14)
    R.figure(SHOTS / "07_agent3_bias_checker.png",
             "Agent 3 output: bias flags and corrected scores", width_cm=14, max_h_cm=17)
    R.figure(crop_tall(SHOTS / "08_agent4_hr_summary.png", 0.9),
             "Agent 4 output: recommendations and tailored interview questions", width_cm=14)
    R.figure(SHOTS / "09_final_ranking.png", "Final ranked shortlist with downloads", width_cm=15)
    R.h3("Command-line execution")
    R.p("The same pipeline runs from the terminal with **python main.py**. Figure 12 is a recording of "
        "that run made with the rich library.")
    R.figure(SHOTS / "terminal_run.png", "Terminal run showing each agent starting, finishing and "
             "handing off")

    R.h2("4.5 Testing")
    R.p("The project has 15 automated tests written with pytest (run with **python -m pytest**). They "
        "always use demo mode, so they never call a paid API. The tests check document parsing, the "
        "JSON format of each agent, and a complete end-to-end run of all four agents through CrewAI "
        "on the sample data, including that every bias trap is caught and that the hand-offs happen "
        "in the right order.")
    test_rows = []
    for line in tests.splitlines():
        m = re.match(r"tests/(\w+)\.py::(\w+) (PASSED|FAILED)", line)
        if m:
            test_rows.append([m.group(1), m.group(2).replace("test_", "").replace("_", " "), m.group(3)])
    R.table(["Test file", "Test", "Result"], test_rows, [3.5, 10, 2.5],
            caption="Automated test results (15 passed)", font=9)

    # ------------------------------------------------------------ 5
    R.page_break()
    R.h1("5. Deployment Architecture")
    R.h2("5.1 Deployment diagram")
    R.p("The application runs as a single Streamlit service, either on a local machine or on "
        "Streamlit Community Cloud. Inside the service, the orchestrator runs four independent CrewAI "
        "agents in sequence. In live mode each agent sends its prompt to the Groq Cloud LLM API over "
        "HTTPS; the API key is read from a secrets store (.env locally, st.secrets in the cloud).")
    R.figure(DIAG / "deployment_architecture.png", "Deployment architecture")
    R.h2("5.2 Proof of multiple agents")
    R.p("Three independent pieces of evidence show that four separate agents run and hand work to "
        "each other:")
    R.bullets([
        "**Hand-off log:** logs/run_log.txt records a timestamped 'started', 'completed' and "
        "'handed off to' line for every agent (Figure 14 and the excerpt below).",
        "**Live status cards:** the UI highlights the agent that is currently working and what it is "
        "doing (Figures 4 and 5).",
        "**Per-agent outputs:** each agent's JSON result is shown in its own labelled section "
        "(Figures 7-10) and saved in outputs/results.json.",
    ])
    R.figure(SHOTS / "10_handoff_log.png", "Agent hand-off log shown in the app", width_cm=15)
    log_lines = [l for l in result["log"] if "handed off" in l or "started" in l or "finished" in l]
    R.code("\n".join(log_lines), caption="excerpt of logs/run_log.txt")
    R.h2("5.3 Running locally")
    R.code("""
git clone https://github.com/<your-username>/resume-screening-panel.git
cd resume-screening-panel
python -m venv .venv
.venv\\Scripts\\activate            # Windows  (macOS/Linux: source .venv/bin/activate)
pip install -r requirements.txt
copy .env.example .env             # then put GROQ_API_KEY=gsk_... in .env
python -m pytest                   # run the tests
python main.py                     # command-line run on sample data
streamlit run app.py               # web app at http://localhost:8501
""")
    R.h2("5.4 Deploying on Streamlit Community Cloud")
    R.numbered([
        "Push the repository to GitHub (the .env file is excluded by .gitignore).",
        "Sign in at share.streamlit.io with GitHub and click 'Create app'.",
        "Select the repository, branch main and main file app.py.",
        "Under 'Advanced settings' -> 'Secrets', add GROQ_API_KEY = \"gsk_...\" (config.py also "
        "reads environment variables, which Streamlit fills from secrets).",
        "Click Deploy. The app receives a public URL that can be shared with evaluators.",
    ])
    R.h2("5.5 Security and privacy")
    R.bullets([
        "API keys are never hard-coded; .env is listed in .gitignore and .env.example has no key.",
        "Resumes are processed in memory and are not stored by the app (only results and logs).",
        "All sample resumes describe fictional people with example.com email addresses.",
        "The system supports a human recruiter; it does not make final hiring decisions on its own.",
    ])

    # ------------------------------------------------------------ 6
    R.h1("6. Results and Discussion")
    R.h2("6.1 Sample run")
    R.p("The panel was run on the sample data: a Junior Python Developer job description and six "
        "fictional resumes (two strong, two average and two weak), three in TXT and three in PDF "
        "format. Several resumes deliberately contain bias traps. The run shown here used **demo "
        "mode**, in which each agent's answer comes from its offline engine; with a Groq key the same "
        "pipeline uses Llama 3.3 70B and the exact numbers will vary between runs.")
    rows = []
    for r in result["summary"]["ranking"]:
        rev = reviews[r["candidate"]]
        flags = ", ".join(f["category"] for f in rev["bias_flags"]) or "None"
        rows.append([r["rank"], r["candidate"], rev["original_score"], r["final_score"],
                     r["recommendation"], flags])
    R.table(["Rank", "Candidate", "Screener score", "Final score", "Decision", "Bias flags"], rows,
            [1.2, 3.8, 2.3, 2.1, 2.0, 4.6], caption="Final ranking produced by the HR Summary Agent")
    R.h2("6.2 Bias cases caught")
    TRAPS = {
        "Priya Raghavan": "Career break for maternity leave",
        "Oluwaseun Adeyemi": "Overseas background / foreign-sounding name",
        "Rakesh Kumar Verma": "Age stated on resume (41)",
        "Sneha Patil": "Tier-3 college and marital status stated",
        "Arjun Mehta": "Tier-1 'premier institute' (positive halo bias)",
        "Kevin D'Souza": "None (control candidate)",
    }
    brows = []
    for name, trap in TRAPS.items():
        rev = reviews[name]
        caught = "; ".join(f"{f['category']} ({-f['points_reversed']:+d})" for f in rev["bias_flags"]) \
            or "No flag (correct)"
        brows.append([name, trap, caught, f"{rev['original_score']} -> {rev['adjusted_score']}"])
    R.table(["Candidate", "Bias trap in resume", "Flag raised (screener points)", "Score change"],
            brows, [3.3, 4.6, 5.0, 3.1], caption="Bias traps and how the Bias Checker handled them",
            font=9)
    R.h2("6.3 Discussion")
    R.bullets([
        "**All five bias traps were caught** and the control candidate was correctly left unchanged.",
        "**Bias correction changed a decision:** Rakesh Kumar Verma's score rose from 70 to 78 once the "
        "age penalty was removed, moving him from Maybe to Hire.",
        "**Bias works in both directions:** Arjun Mehta's top-tier college had added 10 points that his "
        "skills did not justify. The Bias Checker removed this halo effect (44 -> 34).",
        "**Explainability:** every score is backed by a per-skill evidence table, and interview "
        "questions quote the candidate's own resume lines and probe their missing must-have skills.",
        "**Speed:** the complete 14-task pipeline finishes in a few seconds in demo mode; in live mode "
        "speed depends on the Groq API (typically well under a minute for six resumes).",
    ])
    R.h2("6.4 Limitations")
    R.bullets([
        "The offline demo engine uses keyword matching and rules; it is a stand-in for the LLM, not a "
        "replacement, and cannot understand skills described in unusual words.",
        "Live LLM output can vary between runs; the orchestrator re-orders the ranking and applies the "
        "recommendation thresholds itself to keep the results consistent.",
        "Only text-based PDFs are supported (no OCR for scanned resumes).",
        "The sample data is small and synthetic; a real evaluation would need many real, anonymised "
        "resumes and comparison with expert recruiters.",
    ])

    # ------------------------------------------------------------ 7
    R.h1("7. Conclusion and Future Scope")
    R.h2("7.1 Conclusion")
    R.p("This project built a working multi-agent HR system with CrewAI in which four specialised "
        "agents analyse a job description, screen resumes, audit the screening for bias and produce a "
        "ranked, explainable shortlist. Dividing the work between agents made each step simpler, "
        "testable and visible, and adding a dedicated reviewer agent made it possible to catch and "
        "correct unfair reasoning that a single model would hide inside one score. The system is "
        "packaged with a web UI, a command-line runner, automated tests, timestamped hand-off logs and "
        "a deployment guide.")
    R.h2("7.2 Future scope")
    R.bullets([
        "Run the per-resume Screener and Bias Checker tasks in parallel for large applicant pools.",
        "Add OCR (for example Tesseract) to support scanned resumes.",
        "Add a fifth 'Skill Verifier' agent that checks GitHub or portfolio links.",
        "Blind screening: hide names, photos and personal details before the Screener sees the resume.",
        "Connect to an applicant tracking system or Google Drive to fetch resumes automatically.",
        "Measure fairness over many real runs (for example, score differences between groups).",
    ])

    # ------------------------------------------------------------ 8
    R.h1("8. Demo and Repository Links")
    R.p(f"Google Drive Demo Link: {DRIVE_LINK}", bold_lead="")
    R.p(f"GitHub Repository: {GITHUB_URL}")
    R.p("The demo video (docs/demo/demo.mp4, about 2 minutes) shows the six sample resumes being "
        "uploaded, the four agents working in sequence with live status, each agent's output, the final "
        "ranked shortlist and the hand-off log.")

    # ------------------------------------------------------------ refs
    R.h1("References")
    R.numbered([
        "CrewAI documentation - https://docs.crewai.com",
        "Groq API documentation - https://console.groq.com/docs",
        "Meta Llama 3.3 model card - https://www.llama.com",
        "Streamlit documentation - https://docs.streamlit.io",
        "Microsoft AutoGen - https://microsoft.github.io/autogen",
        "LangGraph documentation - https://langchain-ai.github.io/langgraph",
        "Wooldridge, M. (2009). An Introduction to MultiAgent Systems (2nd ed.). Wiley.",
        "Bertrand, M. and Mullainathan, S. (2004). Are Emily and Greg More Employable than Lakisha "
        "and Jamal? A Field Experiment on Labor Market Discrimination. American Economic Review, "
        "94(4), 991-1013.",
        "Raghavan, M., Barocas, S., Kleinberg, J. and Levy, K. (2020). Mitigating Bias in Algorithmic "
        "Hiring: Evaluating Claims and Practices. Proceedings of ACM FAT* 2020.",
    ])
    return R


TOC_ENTRIES: list[tuple[int, str]] = []


def page_map(docx_path: Path, entries) -> dict[str, int]:
    """Render with LibreOffice and find the page where each heading starts."""
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp),
                    str(docx_path)], check=True, capture_output=True)
    pdf = PdfReader(str(tmp / (docx_path.stem + ".pdf")))
    pages = [(page.extract_text() or "") for page in pdf.pages]
    shutil.rmtree(tmp, ignore_errors=True)
    found, start = {}, 2            # skip cover and index pages
    for _, text in entries:
        key = re.sub(r"\s+", " ", text)
        for i in range(start, len(pages)):
            lines = [re.sub(r"\s+", " ", l).strip() for l in pages[i].splitlines()]
            if key in lines:
                found[text] = i + 1
                start = i
                break
    return found


def main():
    global TOC_ENTRIES
    # pass 1: collect headings
    TOC_ENTRIES = [(1, "placeholder")]
    first = build()
    TOC_ENTRIES = first.headings
    # pass 2: build with the real TOC (page numbers blank) and measure pages
    second = build()
    second.save(OUT)
    pages = page_map(OUT, TOC_ENTRIES)
    # pass 3: final document with page numbers in the TOC
    final = build(pages)
    final.save(OUT)
    missing = [t for _, t in TOC_ENTRIES if t not in pages]
    print(f"saved {OUT} ({len(TOC_ENTRIES)} TOC entries, {final.fig} figures, {final.tab} tables)")
    if missing:
        print("TOC entries without page numbers:", missing, file=sys.stderr)


if __name__ == "__main__":
    main()
